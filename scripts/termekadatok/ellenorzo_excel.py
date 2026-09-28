"""Ellenőrző Excel a KÉP NÉLKÜLI termékekről (data/termek_ellenorzes.xlsx).

Csak azok a termékek kerülnek bele, amelyeknek most 0 képe van. Oszlopsorrend: beszállító, megnevezés,
TÖRÖLHETŐ, megjegyzés, majd a többi adat. Termékenként: van-e műszaki adat, honnan jött, hol kerestük és hol nem találtuk,
valamint üres oszlopok a Quadris visszajelzéséhez (TÖRÖLHETŐ, megjegyzés). A „Beszállítók” lap
képletekkel összesít beszállítónként, így a visszaküldött fájlban a jelölések után is friss.

Használat: python3 scripts/termekadatok/ellenorzo_excel.py
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from common import ROOT  # noqa: E402

OUT = ROOT / "data/termek_ellenorzes.xlsx"
SITE = "https://website-9ib.pages.dev"
NAMEBASED = "Quadris adatok"  # nevbol.py: adatlap a Quadris-megnevezésből, nem beszállítói forrás

# Beszállító -> hol kerestük (a scripts/termekadatok/ letöltői alapján)
SEARCHED = {
    "EMEPE": "mpsteel.es – MP Steel katalógus 12/2016 (PDF)",
    "FTS": "fts-farina.it webáruház (WooCommerce API); FTS 2024 katalógus (PDF)",
    "Ital Accessori": "Ital Accessori teljes PDF katalógus (ital-accessori.sk)",
    "Alu-SV": "alu-sv.com webáruház, keresés rendelési számra",
    "ALCOMET": "Alcomet 2017 szabványprofil-katalógus (PDF)",
    "Constellium Decin": "alu-sv.com (a kód utolsó számjegyei); Ital Accessori katalógus; Drive: CONSTELLIUM DECIN Excel – nem letölthető (>10 MB)",
    "ADAICO": "adaico.com webáruház; ADAICO 2025 katalógus (PDF); ADA-Slider vizsgálati jelentés (Drive)",
    "Edscha": "Edscha TS / SESAM alkatrész-katalógus 2024 (edschats.com); Edscha Compact alkatrész-katalógus 2018 (Drive)",
    "Versus": "versus-omega.com termékoldalak; Duo Trike Light alkatrész-PDF (Drive); Versus-Omega online katalógus 2026 (FlippingBook)",
    "JONESCO": "jonesco-plastics.com termékoldalak",
    "G&C termékek": "gnc-systems.com termékoldalak és adatlapok",
    "Sand-Profile": "SAND Profile teljes PDF katalógus",
    "PARLOK": "parlok.com termékoldalak",
    "Pastore": "pastorelombardi.com kereső",
    "Klöckner": "shop.kloeckner.at webáruház",
    "RE-ALL": "re-all.it katalógusfejezetek (PDF) és oldaltérkép; RE-ALL 19 / 21 katalógus (Drive)",
    "POLSER": "polser.com Frigoser oldal és adatlap",
    "BODEGA": "Bodega Bordwände katalógus 2010 (PDF); bodega.it oldaltérkép",
    "Car-Alu": "caralushop.nl webshop",
    "Cimaplast": "cimaplast.it kategóriaoldalak",
    "ESAL": "ESAL Catalogo Sponde 2020 (PDF); ESAL Profili commerciali (PDF)",
    "Metra": "Metra Fahrzeugbau katalógus (Drive)",
    "Grupa Kęty": "Grupa Kęty W2850 gyári rajz (Drive)",
    "Exlabesa PL": "Exlabesa 2023 standard katalógus (Drive)",
    "Exlabesa ES": "Exlabesa 2023 standard katalógus (Drive)",
    "Pommier": "pommier.eu termékoldalak; Pommier 2023–24 katalógus (PDF)",
    "Pommier Furgocar": "pommier.eu termékoldalak; Pommier 2023–24 katalógus (PDF)",
    "BMC Flexpar": "BMC termékoldalak; BMC 2025 katalógus (PDF)",
    "Dost": "Dost Teknik 2024 katalógus (PDF)",
    "Plastic Padana": "plasticpadana.it termékoldalak",
    "COPAR": "copar.it oldaltérkép; Co.Par. Mini Box / Nova Box típuslapok (Drive)",
    "Industrilas": "industrilas.com kereső; Industrilås 68-S1 termékkép (Drive)",
    "IndesCar": "indescar.com webáruház (WooCommerce API)",
    "Takler": "taklergroup.com oldaltérkép bejárása",
    "Profilpol": "profilpolsystem.pl oldaltérkép bejárása",
    "Cargoframes Czech": "cargoframes.eu oldaltérkép; Cargo Frames profilrajzok és műszaki lista (Quadris e-mail)",
}
# ezeknél nincs beszállítói forrás: az adatlap (méretek, elméleti tömeg) a Quadris-megnevezésből készül
NAME_ONLY = {"Peri", "BAYU", "Gummitrading", "VR-Trade", "Novelis", "(nincs megadva)", ""}
# egyedi megjegyzések
NOTES = {
    "202386-u-130-hossztarto": "A Quadris kérésére most kihagyva (csak a Decin Excelben van rajz).",
}

HEAD = ["Beszállító", "Megnevezés", "TÖRÖLHETŐ", "Megjegyzés (Quadris)", "Quadris cikkszám", "Beszállítói kód",
        "Termékcsoport", "Kategória", "Eredet", "Kép (db)", "Műszaki adat (db)", "Leírás", "Állapot", "Kerestük?",
        "Hol kerestük", "Eredmény / hol nem találtuk", "Adatforrás", "Weboldal", "Azonosító (slug)"]
WIDTH = [18, 44, 13, 36, 16, 18, 26, 26, 14, 9, 11, 9, 20, 11, 52, 52, 40, 40, 36]
# rows() belső sorrendje -> HEAD sorrend
ORDER = [4, 1, 16, 17, 0, 5, 2, 3, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 18]

FONT = "Arial"
HFILL = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))


def names():
    cat = json.loads((ROOT / "src/data/catalog.json").read_text())
    ext = json.loads((ROOT / "src/data/bovitett.json").read_text())
    groups = {g["slug"]: ext.get("renameGroups", {}).get(g["slug"], g["name"]) for g in cat["groups"]}
    cats = {(g["slug"], c["slug"]): c["name"] for g in cat["groups"] for c in g["categories"]}
    groups.update({g["slug"]: g["name"] for g in ext.get("groups", [])})
    cats.update({(c["group"], c["slug"]): c["name"] for c in ext.get("categories", [])})
    return cat, ext, groups, cats


def rows():
    cat, ext, groups, cats = names()
    extra = json.loads((ROOT / "src/data/termekadatok.json").read_text())
    out = []
    for p in cat["products"]:
        move = ext.get("moves", {}).get(p["slug"])
        g, c = (move["group"], move["category"]) if move else (p["group"], p["category"])
        e = extra.get(p["slug"], {})
        imgs = e.get("images") or p.get("images") or []
        specs = {**(e.get("specs") or {}), **(p.get("specs") or {})}
        desc = bool(p.get("description") or e.get("description"))
        sup = p["supplier"] or "(nincs megadva)"
        src = e.get("source")
        searched = SEARCHED.get(sup)
        if src and src != NAMEBASED:
            result = f"Megtalálva: {e.get('sourceTitle') or src}" + (f" (kód: {e['matchedCode']})" if e.get("matchedCode") else "")
        elif searched:
            result = f"Nem találtuk: {searched}" + ("; adatlap a Quadris-megnevezésből" if src == NAMEBASED else "")
        elif src == NAMEBASED:
            result = "Beszállítói forrás nincs – adatlap a Quadris-megnevezésből (méretek, elméleti tömeg)"
        else:
            result = "Nem kerestük – nincs ismert weboldal vagy katalógus"
        if p["slug"] in NOTES:
            result += " – " + NOTES[p["slug"]]
        where = searched or ("Quadris-megnevezés (nincs beszállítói forrás)" if sup in NAME_ONLY or src == NAMEBASED else "–")
        out.append([p["code"], p["name"], groups.get(g, g), cats.get((g, c), c), sup, p["supplierCode"], "Quadris Excel",
                    len(imgs), len(specs), "van" if desc else "nincs", None, "Igen" if searched else "Nem", where, result,
                    e.get("sourceUrl") or "", f"{SITE}/termek/{p['slug']}/", "", "", p["slug"]])
    for p in ext.get("products", []):
        sup = p.get("source") or ""
        out.append([p.get("code") or "", p["name"], groups.get(p["group"], p["group"]), cats.get((p["group"], p["category"]), p["category"]),
                    sup, p.get("code") or "", "Teljes kínálat", len(p.get("images") or []), len(p.get("specs") or {}),
                    "van" if p.get("description") else "nincs", None, "Igen", f"{sup} teljes kínálata (gyártói katalógus / weboldal)",
                    "Megtalálva: a beszállító teljes kínálatából felvett termék", p.get("sourceUrl") or "",
                    f"{SITE}/termek/{p['slug']}/", "", "", p["slug"]])
    out = [[r[i] for i in ORDER] for r in out if r[7] == 0]  # csak a kép nélküliek, új oszlopsorrendben
    out.sort(key=lambda r: (r[0].lower(), r[6], r[1]))
    return out


def main():
    data = rows()
    wb = Workbook()

    # --- Útmutató ---
    ws0 = wb.active
    ws0.title = "Útmutató"
    guide = [
        ("Termék-ellenőrző táblázat – Quadris weboldal", True),
        ("", False),
        ("Mit mutat?", True),
        ("A „Termékek” lapon a jelenleg KÉP NÉLKÜLI termékek vannak, soronként egy: van-e műszaki adata, honnan jött az adat, hol kerestük és hol nem találtuk.", False),
        ("A „Beszállítók” lap beszállítónként összesíti a kép nélküli termékeket (képletekkel, a jelölések után is frissül).", False),
        ("", False),
        ("Mit kell kitölteni? (sárga oszlopok a „Termékek” lapon)", True),
        ("C – TÖRÖLHETŐ: legördülőből „IGEN”, ha a terméket le lehet venni a weboldalról és mindenhonnan (katalógus, kereső, ajánlatkérő).", False),
        ("D – Megjegyzés (Quadris): bármi, amit tudnunk kell – pl. „kép: a beszállító új katalógusában a 34. oldalon”, „helyes cikkszám: 7318”, „nem forgalmazzuk”.", False),
        ("Példa: C = IGEN, D = „Kifutott termék, nem rendelhető.”", False),
        ("A többi oszlopot kérjük, ne módosítsák – az S oszlop (azonosító) alapján olvassuk vissza a fájlt.", False),
        ("", False),
        ("Oszlopok jelentése", True),
        ("Kép (db) / Műszaki adat (db): a weboldalon megjelenő képek és adatsorok száma.", False),
        ("Állapot: Rendben / Nincs kép / Nincs adat / Nincs kép és adat (képlet a két számból).", False),
        ("Kerestük?: volt-e ismert beszállítói forrás (weboldal, webáruház, katalógus), ahol a tételt kerestük.", False),
        ("Hol kerestük: a beszállító forrásai, amelyekben a tételt kerestük.", False),
        ("Eredmény: hol találtuk meg (forrás és egyező gyártói kód), vagy hogy a felsorolt helyeken nem találtuk.", False),
        ("„adatlap a Quadris-megnevezésből”: beszállítói adat nincs, a méretek és az elméleti tömeg a megnevezésből számolva.", False),
        ("Eredet: „Quadris Excel” = a Quadris terméklistájából; „Teljes kínálat” = a beszállító teljes kínálatából felvett termék.", False),
        ("", False),
        ("Forrás: a weboldal adatai (src/data/catalog.json, termekadatok.json, bovitett.json) – generálva: scripts/termekadatok/ellenorzo_excel.py", False),
    ]
    for i, (t, bold) in enumerate(guide, 1):
        c = ws0.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 10, color="1F3864" if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 130
    for r in (8, 9):
        ws0.cell(row=r, column=1).fill = INPUT

    # --- Termékek ---
    ws = wb.create_sheet("Termékek")
    ws.append(HEAD)
    for r in data:
        ws.append(r)
    n = len(data) + 1
    for i in range(2, n + 1):
        ws.cell(row=i, column=13, value=f'=IF(AND(J{i}>0,K{i}>0),"Rendben",IF(AND(J{i}=0,K{i}=0),"Nincs kép és adat",IF(J{i}=0,"Nincs kép","Nincs adat")))')
        for col in (17, 18):
            c = ws.cell(row=i, column=col)
            if c.value:
                c.hyperlink = c.value
                c.font = Font(name=FONT, size=9, color="0563C1", underline="single")
    for col, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for c in (ws["C1"], ws["D1"]):
        c.fill = PatternFill("solid", fgColor="BF8F00")
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            if c.hyperlink is None:
                c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, 4, 15, 16))
            c.border = THIN
        row[2].fill = INPUT
        row[3].fill = INPUT
        row[2].alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{n}"
    ws.row_dimensions[1].height = 30
    dv = DataValidation(type="list", formula1='"IGEN"', allow_blank=True, showDropDown=False)
    dv.error, dv.errorTitle = "Csak „IGEN” vagy üres lehet.", "TÖRÖLHETŐ"
    ws.add_data_validation(dv)
    dv.add(f"C2:C{n}")
    red, orange, green = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FCE4D6", "E2EFDA"))
    ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=[f'M2="Nincs kép és adat"'], fill=red))
    ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=[f'OR(M2="Nincs kép",M2="Nincs adat")'], fill=orange))
    ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=[f'M2="Rendben"'], fill=green))
    ws.conditional_formatting.add(f"A2:B{n}", FormulaRule(formula=[f'$C2="IGEN"'], font=Font(strike=True, color="808080")))
    ws.conditional_formatting.add(f"E2:R{n}", FormulaRule(formula=[f'$C2="IGEN"'], font=Font(strike=True, color="808080")))

    # --- Beszállítók ---
    wsb = wb.create_sheet("Beszállítók", 1)
    bh = ["Beszállító", "Kép nélküli termékek", "Ebből műszaki adat sincs", "Ebből van műszaki adat",
          "Kerestük?", "Hol kerestük", "Törlésre jelölve"]
    wsb.append(bh)
    sups = []
    for r in data:
        if r[0] not in sups:
            sups.append(r[0])
    rng = lambda col: f"Termékek!${col}$2:${col}${n}"  # noqa: E731
    for i, s in enumerate(sups, 2):
        first = next(r for r in data if r[0] == s)
        wsb.append([s,
                    f'=COUNTIF({rng("A")},A{i})',
                    f'=COUNTIFS({rng("A")},A{i},{rng("M")},"Nincs kép és adat")',
                    f'=B{i}-C{i}',
                    first[13], first[14],
                    f'=COUNTIFS({rng("A")},A{i},{rng("C")},"IGEN")'])
    last = len(sups) + 1
    tot = last + 1
    wsb.cell(row=tot, column=1, value="Összesen")
    for col in "BCDG":
        wsb[f"{col}{tot}"] = f"=SUM({col}2:{col}{last})"
    for c in wsb[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in wsb.iter_rows(min_row=2, max_row=tot):
        for c in row:
            c.font = Font(name=FONT, size=9, bold=c.row == tot)
            c.alignment = Alignment(vertical="top", wrap_text=c.column == 6)
            c.border = THIN
    for col, w in zip("ABCDEFG", (22, 12, 14, 14, 10, 80, 12)):
        wsb.column_dimensions[col].width = w
    wsb.freeze_panes = "B2"
    wsb.auto_filter.ref = f"A1:G{last}"
    wsb.conditional_formatting.add(f"C2:C{last}", FormulaRule(formula=["C2>0"], fill=orange))

    wb.calculation.fullCalcOnLoad = True  # a képleteket az Excel megnyitáskor számolja ki
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: {len(data)} termék, {len(sups)} beszállító")


if __name__ == "__main__":
    main()
